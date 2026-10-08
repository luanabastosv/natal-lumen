"""A planilha consolidada de uma edicao: o sistema inteiro, num arquivo.

Existe para o dia em que o sistema cair. Toda noite uma copia nova sobe para o
Drive (ver gerar_planilha.py), e a equipe continua acompanhando por ela — e por
isso ela imita a "LISTA GERAL" que a coordenacao mantinha a mao: uma aba por
instituicao, um consolidado e os calculos. Quem conhecia a planilha antiga se
acha nesta sem aprender nada.

E so leitura: uma foto do banco na hora em que foi gerada. Mexer nela nao muda
nada no sistema.

Ficam de fora, de proposito: senhas, hashes, log de atividade e o sorrisometro
(a meta de 1450 nunca entrou no sistema — ver a memoria do painel).
"""

from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.models import (
    Apadrinhamento,
    Autorizacao,
    Cartao,
    Compra,
    Crianca,
    DiaEvento,
    Edicao,
    InstituicaoDia,
    Kit,
    Padrinho,
    Pagamento,
    Recebimento,
    UsuarioEdicao,
)
from app.seeds.perfis_permissoes import PERFIL_COMISSARIO

FUSO = ZoneInfo("America/Sao_Paulo")

NAVY = "153377"
CREME = "FBF7EE"
AMARELO = "FFE7A3"
VERMELHO_CLARO = "FDECEC"

CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor=NAVY)
FUNDO_GRUPO = PatternFill("solid", fgColor=CREME)
FUNDO_DESISTENTE = PatternFill("solid", fgColor=VERMELHO_CLARO)
FUNDO_DESTAQUE = PatternFill("solid", fgColor=AMARELO)
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=13, color=NAVY)
FINA = Side(style="thin", color="D0D5E0")
GRADE = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
CENTRO = Alignment(horizontal="center", vertical="center")

# As colunas de um tipo (cesta ou festa). Iguais nos dois, lado a lado, como na
# planilha antiga.
COLUNAS_TIPO = [
    ("Padrinho", 26),
    ("Situação", 18),
    ("Data pag.", 11),
    ("Valor", 10),
    ("Forma pag.", 12),
    ("Comprovante", 14),
    ("OBS", 22),
    ("Comissário", 18),
    ("Grupo", 12),
]

COLUNAS_CRIANCA = [
    ("Saiu?", 7),
    ("Código", 9),
    ("Nome", 32),
    ("Idade", 7),
    ("Sexo", 6),
]

COLUNAS_FIM = [
    ("Cartão cesta", 12),
    ("Cartão festa", 12),
    ("Autorização", 12),
    ("Kit", 10),
    ("Check-in", 16),
]

SEM = "SEM APADRINHAMENTO"


def _data(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _momento(m) -> str:
    return m.astimezone(FUSO).strftime("%d/%m/%Y %H:%M") if m else ""


def _dinheiro(celula, valor) -> None:
    celula.value = float(valor) if valor is not None else None
    celula.number_format = '"R$" #,##0.00'


class _Dados:
    """Tudo o que a planilha precisa, carregado de uma vez.

    Poucas consultas grandes, e nao uma por crianca: uma edicao tem 1500
    criancas, e a planilha roda de madrugada junto com o backup.
    """

    def __init__(self, db: Session, edicao: Edicao):
        self.edicao = edicao

        self.criancas = db.scalars(
            select(Crianca)
            .where(Crianca.edicao_id == edicao.id)
            .options(
                joinedload(Crianca.instituicao),
                joinedload(Crianca.dia_evento),
                joinedload(Crianca.comissario),
            )
            .order_by(Crianca.codigo)
        ).all()
        ids = [c.id for c in self.criancas] or [0]

        self.apadrinhamento: dict[tuple[int, str], Apadrinhamento] = {
            (a.crianca_id, a.tipo): a
            for a in db.scalars(
                select(Apadrinhamento)
                .where(Apadrinhamento.crianca_id.in_(ids))
                .options(
                    joinedload(Apadrinhamento.padrinho),
                    joinedload(Apadrinhamento.pagamento),
                    joinedload(Apadrinhamento.comissario),
                )
            ).all()
        }

        self.cartoes = {
            (c, t) for c, t in db.execute(
                select(Cartao.crianca_id, Cartao.tipo).where(Cartao.crianca_id.in_(ids))
            ).all()
        }
        self.autorizacoes = set(
            db.scalars(select(Autorizacao.crianca_id).where(Autorizacao.crianca_id.in_(ids))).all()
        )
        self.kits = dict(
            db.execute(select(Kit.crianca_id, Kit.status).where(Kit.crianca_id.in_(ids))).all()
        )

        # O time da edicao, com o grupo de cada um NESTA edicao
        # (usuario_edicao.grupo_id) e o perfil — e o perfil que diz quem e
        # comissario e entra no controle por comissario mesmo sem crianca.
        self.time = db.scalars(
            select(UsuarioEdicao)
            .where(UsuarioEdicao.edicao_id == edicao.id, UsuarioEdicao.ativo.is_(True))
            .options(
                joinedload(UsuarioEdicao.grupo),
                joinedload(UsuarioEdicao.usuario),
                joinedload(UsuarioEdicao.perfil),
            )
        ).all()
        self.grupo_do_usuario = {v.usuario_id: v.grupo.nome for v in self.time if v.grupo}

        self.dias = db.scalars(
            select(DiaEvento).where(DiaEvento.edicao_id == edicao.id).order_by(DiaEvento.data)
        ).all()
        self.onibus = {
            i: o
            for i, o in db.execute(
                select(InstituicaoDia.instituicao_id, InstituicaoDia.onibus).where(
                    InstituicaoDia.edicao_id == edicao.id
                )
            ).all()
        }

        self.padrinhos = db.scalars(
            select(Padrinho)
            .where(Padrinho.edicao_id == edicao.id)
            .options(selectinload(Padrinho.apadrinhamentos).joinedload(Apadrinhamento.crianca))
            .order_by(Padrinho.nome)
        ).all()
        self.pagamentos = db.scalars(
            select(Pagamento)
            .join(Padrinho, Padrinho.id == Pagamento.padrinho_id)
            .where(Padrinho.edicao_id == edicao.id)
            .options(
                joinedload(Pagamento.padrinho),
                selectinload(Pagamento.apadrinhamentos).joinedload(Apadrinhamento.crianca),
            )
            .order_by(Pagamento.data)
        ).all()
        self.recebimentos = db.scalars(
            select(Recebimento).where(Recebimento.edicao_id == edicao.id).order_by(Recebimento.data)
        ).all()
        self.compras = db.scalars(
            select(Compra).where(Compra.edicao_id == edicao.id).order_by(Compra.data)
        ).all()

    def instituicoes(self):
        """(instituicao, dia, criancas), pela ordem do dia e depois do nome —
        a mesma ordem das abas da planilha antiga."""
        grupos = defaultdict(list)
        for c in self.criancas:
            grupos[c.instituicao_id].append(c)
        chave = lambda cs: (  # noqa: E731
            cs[0].dia_evento.data if cs[0].dia_evento else datetime.max.date(),
            cs[0].instituicao.nome,
        )
        for cs in sorted(grupos.values(), key=chave):
            yield cs[0].instituicao, cs[0].dia_evento, cs


# ------------------------------------------------------------------ montagem


def _cabecalho(ws: Worksheet, linha: int, colunas: list[tuple[str, int]], inicio: int = 1):
    for i, (rotulo, largura) in enumerate(colunas):
        celula = ws.cell(row=linha, column=inicio + i, value=rotulo)
        celula.font = CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.alignment = CENTRO
        celula.border = GRADE
        ws.column_dimensions[get_column_letter(inicio + i)].width = largura


def _valores_do_tipo(dados: _Dados, crianca: Crianca, tipo: str) -> list:
    a = dados.apadrinhamento.get((crianca.id, tipo))
    if a is None:
        return [SEM, "", "", None, "", "", "", "", ""]

    pag = a.pagamento
    comissario = a.comissario or crianca.comissario
    return [
        a.padrinho.nome,
        "Pago" if pag else "Promessa (falta pagar)",
        _data(pag.data) if pag else "",
        a.valor,
        (pag.forma or "") if pag else "",
        (pag.comprovante_drive_link or ("no sistema" if pag.comprovante_arquivo else ""))
        if pag
        else "",
        (pag.observacoes or "") if pag else "",
        comissario.nome if comissario else "",
        dados.grupo_do_usuario.get(comissario.id, "") if comissario else "",
    ]


def _linha_crianca(dados: _Dados, c: Crianca) -> list:
    return [
        "SIM" if c.desistiu_em else "",
        c.codigo,
        c.nome,
        c.idade,
        c.sexo,
        *_valores_do_tipo(dados, c, "cesta"),
        *_valores_do_tipo(dados, c, "festa"),
        "Subiu" if (c.id, "cesta") in dados.cartoes else "",
        "Subiu" if (c.id, "festa") in dados.cartoes else "",
        "Subiu" if c.id in dados.autorizacoes else "",
        "Montado" if dados.kits.get(c.id) == "montado" else "",
        _momento(c.checkin_em),
    ]


def _lista_de_criancas(ws: Worksheet, dados: _Dados, criancas, titulo: str, extras=()):
    """A tabela de criancas: uma aba de instituicao, ou o consolidado.

    `extras` sao colunas a mais no comeco (o consolidado ganha instituicao e
    dia, que numa aba de instituicao seriam a mesma palavra em toda linha).
    """
    base = [*((rotulo, largura) for rotulo, largura, _ in extras), *COLUNAS_CRIANCA]
    colunas = [*base, *COLUNAS_TIPO, *COLUNAS_TIPO, *COLUNAS_FIM]

    ws["A1"] = titulo
    ws["A1"].font = TITULO

    # Linha 2: os blocos Cesta e Festa, como na planilha antiga.
    for rotulo, inicio in (("CESTA", len(base) + 1), ("FESTA", len(base) + len(COLUNAS_TIPO) + 1)):
        fim = inicio + len(COLUNAS_TIPO) - 1
        ws.merge_cells(start_row=2, start_column=inicio, end_row=2, end_column=fim)
        celula = ws.cell(row=2, column=inicio, value=rotulo)
        celula.font = NEGRITO
        celula.fill = FUNDO_GRUPO
        celula.alignment = CENTRO

    _cabecalho(ws, 3, colunas)

    col_valor = [len(base) + 4, len(base) + len(COLUNAS_TIPO) + 4]
    for n, c in enumerate(criancas, start=4):
        valores = [*(valor(c) for _, _, valor in extras), *_linha_crianca(dados, c)]
        for i, v in enumerate(valores, start=1):
            celula = ws.cell(row=n, column=i)
            if i in col_valor:
                _dinheiro(celula, v)
            else:
                celula.value = v
            celula.border = GRADE
            if c.desistiu_em:
                celula.fill = FUNDO_DESISTENTE

    ws.freeze_panes = ws.cell(row=4, column=len(base) + 1)
    if criancas:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(colunas))}{3 + len(criancas)}"
    return len(colunas)


def _idade_por_sexo(ws: Worksheet, criancas, coluna: int, linha: int = 3):
    """O quadro "quant. homens e mulheres por idade", ao lado da lista."""
    contagem = Counter((c.idade, c.sexo) for c in criancas if not c.desistiu_em)
    if ws.cell(row=linha - 1, column=coluna).value is None:
        ws.cell(row=linha - 1, column=coluna, value="POR IDADE E SEXO").font = NEGRITO
    _cabecalho(ws, linha, [("Idade", 8), ("Masc", 7), ("Fem", 7)], inicio=coluna)
    idades = sorted({i for i, _ in contagem})
    for n, idade in enumerate(idades, start=linha + 1):
        for k, v in enumerate((idade, contagem[(idade, "M")], contagem[(idade, "F")])):
            celula = ws.cell(row=n, column=coluna + k, value=v)
            celula.border = GRADE
            celula.alignment = CENTRO


def _calculos(ws: Worksheet, dados: _Dados):
    """A aba de numeros: por dia, por instituicao, e o total da campanha."""
    ws["A1"] = f"{dados.edicao.nome} — gerada em {datetime.now(FUSO):%d/%m/%Y %H:%M}"
    ws["A1"].font = TITULO
    ws["A2"] = (
        "Foto do sistema na hora em que foi gerada. Mexer aqui não muda nada no sistema. "
        "Apadrinhamento só conta quando está pago; promessa fica de fora."
    )

    linha = 4
    colunas = [
        ("Instituição", 34),
        ("Ônibus", 9),
        ("Crianças", 10),
        ("Com cesta", 11),
        ("Com festa", 11),
        ("Completas", 11),
        ("Sem padrinho", 13),
        ("% completas", 13),
    ]

    def pago(c, tipo):
        a = dados.apadrinhamento.get((c.id, tipo))
        return a is not None and a.pagamento_id is not None

    por_dia = defaultdict(list)
    for inst, dia, criancas in dados.instituicoes():
        por_dia[dia.id if dia else None].append((inst, criancas))

    nomes_dos_dias = {d.id: d for d in dados.dias}
    totais_gerais = [0] * 6
    for dia_id, linhas in por_dia.items():
        dia = nomes_dos_dias.get(dia_id)
        rotulo = (
            f"{(dia.descricao or '').upper()} {_data(dia.data)}".strip() if dia else "SEM DIA DEFINIDO"
        )
        ws.cell(row=linha, column=1, value=rotulo).font = TITULO
        linha += 1
        _cabecalho(ws, linha, colunas)
        linha += 1
        soma = [0] * 6
        for inst, criancas in linhas:
            vem = [c for c in criancas if not c.desistiu_em]
            cesta = sum(pago(c, "cesta") for c in vem)
            festa = sum(pago(c, "festa") for c in vem)
            completas = sum(pago(c, "cesta") and pago(c, "festa") for c in vem)
            sem = sum(not pago(c, "cesta") and not pago(c, "festa") for c in vem)
            numeros = [dados.onibus.get(inst.id, 0), len(vem), cesta, festa, completas, sem]
            soma = [a + b for a, b in zip(soma, numeros)]
            _linha_de_numeros(ws, linha, inst.nome, numeros)
            linha += 1
        _linha_de_numeros(ws, linha, "TOTAL", soma, negrito=True)
        totais_gerais = [a + b for a, b in zip(totais_gerais, soma)]
        linha += 2

    _linha_de_numeros(ws, linha, "TOTAL GERAL", totais_gerais, negrito=True)
    linha += 3

    # Apadrinhamentos e dinheiro. Desistente fica de fora da conta do que
    # falta (ninguem vai procurar padrinho para quem nao vem), mas o dinheiro
    # que ja entrou por ela continua contado: ele entrou.
    pagos = [a for a in dados.apadrinhamento.values() if a.pagamento_id is not None]
    vem = [c for c in dados.criancas if not c.desistiu_em]
    faltam = sum(not pago(c, t) for c in vem for t in ("cesta", "festa"))
    vagas = len(vem) * 2
    ws.cell(row=linha, column=1, value="APADRINHAMENTOS").font = TITULO
    linha += 1
    _cabecalho(
        ws,
        linha,
        [("Apadrinhados (pagos)", 34), ("Valor total", 9), ("Faltam", 10), ("% feito", 11)],
    )
    linha += 1
    valores = [len(pagos), sum((a.valor for a in pagos), Decimal(0)), faltam,
               (vagas - faltam) / vagas if vagas else 0]
    for k, v in enumerate(valores, start=1):
        celula = ws.cell(row=linha, column=k)
        if k == 2:
            _dinheiro(celula, v)
        elif k == 4:
            celula.value = v
            celula.number_format = "0.0%"
            celula.fill = FUNDO_DESTAQUE
        else:
            celula.value = v
        celula.border = GRADE
        celula.alignment = CENTRO
    ws.column_dimensions["B"].width = 14
    linha += 3

    # Formas de pagamento: quantas e quanto, de padrinho e de doacao.
    formas = defaultdict(lambda: [0, Decimal(0)])
    for p in dados.pagamentos:
        formas[p.forma or "Não informada"][0] += 1
        formas[p.forma or "Não informada"][1] += p.valor
    for r in dados.recebimentos:
        formas[r.forma or "Não informada"][0] += 1
        formas[r.forma or "Não informada"][1] += r.valor
    ws.cell(row=linha, column=1, value="FORMAS DE PAGAMENTO").font = TITULO
    linha += 1
    _cabecalho(ws, linha, [("Forma", 34), ("Quantos", 9), ("Valor", 10)])
    linha += 1
    for forma, (quantos, valor) in sorted(formas.items()):
        ws.cell(row=linha, column=1, value=forma).border = GRADE
        ws.cell(row=linha, column=2, value=quantos).border = GRADE
        celula = ws.cell(row=linha, column=3)
        _dinheiro(celula, valor)
        celula.border = GRADE
        linha += 1


def _linha_de_numeros(ws, linha, rotulo, numeros, negrito=False):
    onibus, criancas, cesta, festa, completas, sem = numeros
    valores = [rotulo, onibus, criancas, cesta, festa, completas, sem,
               completas / criancas if criancas else None]
    for k, v in enumerate(valores, start=1):
        celula = ws.cell(row=linha, column=k, value=v)
        celula.border = GRADE
        if k > 1:
            celula.alignment = CENTRO
        if k == 8:
            celula.number_format = "0.0%"
        if negrito:
            celula.font = NEGRITO


def _pago(dados: _Dados, c: Crianca, tipo: str) -> bool:
    a = dados.apadrinhamento.get((c.id, tipo))
    return a is not None and a.pagamento_id is not None


def _controles(ws: Worksheet, dados: _Dados):
    """A aba CONTROLES da planilha antiga: quem saiu, o controle por
    comissario e o de idade.

    O bloco "Forms x consolidado & ajuste dos duplicados" nao vem: ele
    conferia o Google Forms contra a planilha, e no sistema apadrinhamento
    duplicado e impossivel — cada crianca tem no maximo uma cesta e uma festa.
    """
    vem = [c for c in dados.criancas if not c.desistiu_em]
    sairam = len(dados.criancas) - len(vem)

    ws["A1"] = "SAÍRAM"
    ws["A1"].font = TITULO
    ws["B1"] = sairam
    ws["B1"].font = NEGRITO
    ws["C1"] = "crianças desistiram e ficam fora das contas abaixo"

    ws.cell(row=3, column=1, value="CONTROLE POR COMISSÁRIO").font = TITULO
    colunas = [
        ("Comissário", 30), ("Grupo", 16), ("Crianças", 10), ("Com cesta", 11),
        ("Com festa", 11), ("Completas", 11), ("Falta apadrinhar", 16), ("% falta", 10),
    ]
    _cabecalho(ws, 4, colunas)

    por_responsavel = defaultdict(list)
    for c in vem:
        por_responsavel[c.comissario_id].append(c)

    # O time manda, e nao as criancas: o comissario sem nenhuma crianca na mao
    # e justamente o que a coordenacao precisa enxergar. Mesma regra do painel.
    nomes = {v.usuario_id: v.usuario.nome for v in dados.time}
    responsaveis = {v.usuario_id for v in dados.time if v.perfil.nome == PERFIL_COMISSARIO}
    responsaveis |= {i for i in por_responsavel if i is not None}
    for c in vem:
        if c.comissario_id and c.comissario_id not in nomes:
            nomes[c.comissario_id] = c.comissario.nome

    ordem = sorted(
        responsaveis,
        key=lambda i: (i not in dados.grupo_do_usuario, dados.grupo_do_usuario.get(i, ""), nomes[i]),
    )
    if None in por_responsavel:
        ordem.append(None)

    linha = 5
    for i in ordem:
        criancas = por_responsavel.get(i, [])
        completas = sum(_pago(dados, c, "cesta") and _pago(dados, c, "festa") for c in criancas)
        valores = [
            nomes[i] if i else "Sem comissário",
            dados.grupo_do_usuario.get(i, "") if i else "",
            len(criancas),
            sum(_pago(dados, c, "cesta") for c in criancas),
            sum(_pago(dados, c, "festa") for c in criancas),
            completas,
            len(criancas) - completas,
            (len(criancas) - completas) / len(criancas) if criancas else None,
        ]
        for k, v in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=k, value=v)
            celula.border = GRADE
            if k > 2:
                celula.alignment = CENTRO
            if k == 8:
                celula.number_format = "0.0%"
        linha += 1
    ws.freeze_panes = "B5"

    linha += 2
    ws.cell(row=linha, column=1, value="CONTROLE DE IDADE (FEMININO X MASCULINO)").font = TITULO
    _idade_por_sexo(ws, dados.criancas, coluna=1, linha=linha + 1)


def _padrinhos(ws: Worksheet, dados: _Dados):
    colunas = [
        ("Nome", 30), ("WhatsApp", 16), ("Email", 28), ("Crianças", 40),
        ("Cestas", 8), ("Festas", 8), ("Combinado", 12), ("Pago", 12), ("Observações", 30),
    ]
    _cabecalho(ws, 1, colunas)
    for n, p in enumerate(dados.padrinhos, start=2):
        aps = p.apadrinhamentos
        valores = [
            p.nome,
            p.whatsapp or "",
            p.email or "",
            ", ".join(
                f"{a.crianca.codigo} {a.tipo}{'' if a.pagamento_id else ' (promessa)'}"
                for a in sorted(aps, key=lambda a: a.crianca.codigo)
            ),
            sum(a.tipo == "cesta" for a in aps),
            sum(a.tipo == "festa" for a in aps),
            sum((a.valor for a in aps), Decimal(0)),
            sum((a.valor for a in aps if a.pagamento_id), Decimal(0)),
            p.observacoes or "",
        ]
        for k, v in enumerate(valores, start=1):
            celula = ws.cell(row=n, column=k)
            if k in (7, 8):
                _dinheiro(celula, v)
            else:
                celula.value = v
            celula.border = GRADE
    ws.freeze_panes = "B2"
    if dados.padrinhos:
        ws.auto_filter.ref = f"A1:I{1 + len(dados.padrinhos)}"


def _financeiro(ws: Worksheet, dados: _Dados):
    """Entradas (padrinhos e doacoes) e saidas, uma lista embaixo da outra."""
    ws["A1"] = "ENTRADAS"
    ws["A1"].font = TITULO
    colunas = [
        ("Data", 11), ("Origem", 14), ("Quem", 30), ("Descrição", 40), ("Valor", 12),
        ("Forma", 12), ("Conferido", 10), ("Comprovante", 30), ("Observações", 30),
    ]
    _cabecalho(ws, 2, colunas)
    linhas = []
    for p in dados.pagamentos:
        criancas = ", ".join(
            f"{a.crianca.codigo} {a.tipo}" for a in sorted(p.apadrinhamentos, key=lambda a: a.crianca.codigo)
        )
        linhas.append((p.data, "Padrinho", p.padrinho.nome, criancas or "sem destino", p.valor,
                       p.forma, p.conferido, p.comprovante_drive_link, p.observacoes))
    for r in dados.recebimentos:
        linhas.append((r.data, r.categoria.capitalize(), r.doador, r.descricao, r.valor,
                       r.forma, r.conferido, r.comprovante_drive_link, r.observacoes))
    linhas.sort(key=lambda x: x[0])
    n = 3
    for data, origem, quem, desc, valor, forma, conferido, link, obs in linhas:
        valores = [_data(data), origem, quem or "", desc, valor, forma or "",
                   "Sim" if conferido else "", link or "", obs or ""]
        for k, v in enumerate(valores, start=1):
            celula = ws.cell(row=n, column=k)
            if k == 5:
                _dinheiro(celula, v)
            else:
                celula.value = v
            celula.border = GRADE
        n += 1
    ws.cell(row=n, column=4, value="TOTAL DAS ENTRADAS").font = NEGRITO
    _dinheiro(ws.cell(row=n, column=5), sum((x[4] for x in linhas), Decimal(0)))
    ws.cell(row=n, column=5).font = NEGRITO

    n += 3
    ws.cell(row=n, column=1, value="SAÍDAS").font = TITULO
    n += 1
    _cabecalho(ws, n, [("Data", 11), ("Categoria", 14), ("Fornecedor", 30), ("Descrição", 40),
                       ("Valor", 12), ("Quantidade", 12)])
    n += 1
    for c in dados.compras:
        valores = [_data(c.data), c.categoria or "", c.fornecedor or "", c.descricao,
                   c.valor_total, c.quantidade]
        for k, v in enumerate(valores, start=1):
            celula = ws.cell(row=n, column=k)
            if k == 5:
                _dinheiro(celula, v)
            else:
                celula.value = v
            celula.border = GRADE
        n += 1
    ws.cell(row=n, column=4, value="TOTAL DAS SAÍDAS").font = NEGRITO
    _dinheiro(ws.cell(row=n, column=5), sum((c.valor_total for c in dados.compras), Decimal(0)))
    ws.cell(row=n, column=5).font = NEGRITO


def _nome_de_aba(nome: str, usados: set[str]) -> str:
    """O Excel e o Google limitam a 31 caracteres e proibem []:*?/\\ no nome."""
    limpo = "".join(ch for ch in nome if ch not in "[]:*?/\\").strip()[:31] or "Instituição"
    final, k = limpo, 2
    while final.upper() in usados:
        sufixo = f" ({k})"
        final = limpo[: 31 - len(sufixo)] + sufixo
        k += 1
    usados.add(final.upper())
    return final


def montar(db: Session, edicao: Edicao) -> bytes:
    """O .xlsx da edicao inteira."""
    dados = _Dados(db, edicao)
    livro = Workbook()
    usados = {"CÁLCULOS", "CONSOLIDADO", "CONTROLES", "PADRINHOS", "FINANCEIRO"}

    # Os calculos vem primeiro: no dia em que o sistema cair, a primeira
    # pergunta e "como estamos", e nao "quem e a crianca MA123".
    calc = livro.active
    calc.title = "CÁLCULOS"
    _calculos(calc, dados)

    consolidado = livro.create_sheet("CONSOLIDADO")
    extras = (
        ("Instituição", 28, lambda c: c.instituicao.nome),
        ("Dia", 12, lambda c: (c.dia_evento.descricao or _data(c.dia_evento.data)) if c.dia_evento else ""),
    )
    _lista_de_criancas(
        consolidado, dados, dados.criancas, f"{edicao.nome} — todas as crianças", extras=extras
    )

    for inst, dia, criancas in dados.instituicoes():
        ws = livro.create_sheet(_nome_de_aba(inst.nome.upper(), usados))
        quando = f" — {(dia.descricao or '').upper()} {_data(dia.data)}".rstrip() if dia else " — SEM DIA"
        largura = _lista_de_criancas(ws, dados, criancas, f"{inst.nome.upper()}{quando}")
        _idade_por_sexo(ws, criancas, coluna=largura + 2)

    _controles(livro.create_sheet("CONTROLES"), dados)
    _padrinhos(livro.create_sheet("PADRINHOS"), dados)
    _financeiro(livro.create_sheet("FINANCEIRO"), dados)

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_do_arquivo(edicao: Edicao, quando: datetime | None = None) -> str:
    quando = quando or datetime.now(FUSO)
    return f"{edicao.nome} — {quando:%Y-%m-%d}"
