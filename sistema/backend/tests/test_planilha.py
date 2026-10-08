"""Testa a planilha consolidada que sobe para o Drive toda noite.

O que ela promete:

  1. uma aba por instituicao, alem de CALCULOS, CONSOLIDADO, PADRINHOS e
     FINANCEIRO;
  2. cada crianca com o padrinho de cesta e de festa, dizendo se e pago ou
     promessa, e "SEM APADRINHAMENTO" quando nao ha;
  3. a desistente marcada em "Saiu?", e fora das contas do que falta;
  4. so o pago conta como apadrinhado nos calculos.

Rodar com:  python -m tests.test_planilha
"""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from app.database import SessionLocal
from app.models import (
    Apadrinhamento,
    Cartao,
    Cidade,
    Crianca,
    DiaEvento,
    Edicao,
    Instituicao,
    InstituicaoDia,
    Padrinho,
    Pagamento,
)
from app.servicos import planilha

MARCA = "ZZ_PLA"

ok = 0
falhas: list[str] = []


def verifica(d: str, c: bool, extra: str = "") -> None:
    global ok
    if c:
        ok += 1
        print(f"  ok    {d}")
    else:
        falhas.append(d)
        print(f"  FALHA {d} {extra}")


def limpar(db) -> None:
    db.rollback()
    cids = [c.id for c in db.scalars(select(Cidade).where(Cidade.nome.ilike(f"{MARCA}%"))).all()]
    if cids:
        eds = list(db.scalars(select(Edicao.id).where(Edicao.cidade_id.in_(cids))).all())
        if eds:
            cris = list(db.scalars(select(Crianca.id).where(Crianca.edicao_id.in_(eds))).all())
            pads = list(db.scalars(select(Padrinho.id).where(Padrinho.edicao_id.in_(eds))).all())
            if cris:
                db.execute(delete(Cartao).where(Cartao.crianca_id.in_(cris)))
                db.execute(delete(Apadrinhamento).where(Apadrinhamento.crianca_id.in_(cris)))
            if pads:
                db.execute(delete(Pagamento).where(Pagamento.padrinho_id.in_(pads)))
                db.execute(delete(Padrinho).where(Padrinho.id.in_(pads)))
            db.execute(delete(Crianca).where(Crianca.edicao_id.in_(eds)))
            db.execute(delete(InstituicaoDia).where(InstituicaoDia.edicao_id.in_(eds)))
            db.execute(delete(DiaEvento).where(DiaEvento.edicao_id.in_(eds)))
            db.execute(delete(Edicao).where(Edicao.id.in_(eds)))
        db.execute(delete(Instituicao).where(Instituicao.cidade_id.in_(cids)))
        db.execute(delete(Cidade).where(Cidade.id.in_(cids)))
    db.commit()


def main() -> None:
    db = SessionLocal()
    limpar(db)

    cidade = Cidade(nome=f"{MARCA} Cidade", uf="CE"); db.add(cidade); db.flush()
    edicao = Edicao(cidade_id=cidade.id, ano=2026, nome=f"{MARCA} Edicao",
                    valor_cesta=130, valor_festa=60); db.add(edicao); db.flush()
    sabado = DiaEvento(edicao_id=edicao.id, data=date(2026, 12, 5), descricao="Sábado")
    db.add(sabado); db.flush()
    escola = Instituicao(cidade_id=cidade.id, nome=f"{MARCA} Escola Esperança")
    sem_dia = Instituicao(cidade_id=cidade.id, nome=f"{MARCA} Creche Sol")
    db.add_all([escola, sem_dia]); db.flush()
    db.add(InstituicaoDia(edicao_id=edicao.id, instituicao_id=escola.id,
                          dia_evento_id=sabado.id, onibus=2))

    def crianca(codigo, inst, dia, idade=8, sexo="F"):
        c = Crianca(edicao_id=edicao.id, instituicao_id=inst.id, dia_evento_id=dia.id if dia else None,
                    codigo=codigo, nome=f"Crianca {codigo}", idade=idade, sexo=sexo)
        db.add(c); db.flush()
        return c

    a1 = crianca("E1", escola, sabado)
    a2 = crianca("E2", escola, sabado, sexo="M")
    a3 = crianca("E3", escola, sabado)
    s1 = crianca("S1", sem_dia, None)
    a3.desistiu_em = func.now()

    padrinho = Padrinho(edicao_id=edicao.id, nome=f"{MARCA} Padrinho Um", whatsapp="85999990000")
    db.add(padrinho); db.flush()
    pag = Pagamento(padrinho_id=padrinho.id, valor=130, data=date(2026, 10, 6), forma="Pix")
    db.add(pag); db.flush()
    db.add(Apadrinhamento(crianca_id=a1.id, padrinho_id=padrinho.id, tipo="cesta",
                          valor=130, pagamento_id=pag.id))
    db.add(Apadrinhamento(crianca_id=a1.id, padrinho_id=padrinho.id, tipo="festa", valor=60))
    db.add(Cartao(crianca_id=a1.id, tipo="cesta", arquivo="zz/E1_cesta.jpg"))
    db.commit()

    try:
        livro = load_workbook(BytesIO(planilha.montar(db, edicao)))

        print("\nAs abas")
        nomes = livro.sheetnames
        verifica("calculos primeiro, depois o consolidado", nomes[:2] == ["CÁLCULOS", "CONSOLIDADO"])
        verifica("uma aba por instituicao",
                 f"{MARCA} ESCOLA ESPERANÇA"[:31] in nomes and f"{MARCA} CRECHE SOL" in nomes, str(nomes))
        verifica("padrinhos e financeiro no fim", nomes[-2:] == ["PADRINHOS", "FINANCEIRO"])

        print("\nA aba da instituicao")
        ws = livro[f"{MARCA} ESCOLA ESPERANÇA"[:31]]
        cabecalho = [c.value for c in ws[3]]
        linhas = {r[1]: r for r in ws.iter_rows(min_row=4, values_only=True) if r[1]}
        col = {nome: cabecalho.index(nome) for nome in ("Saiu?", "Padrinho", "Situação", "Forma pag.")}
        festa = cabecalho.index("Padrinho", col["Padrinho"] + 1)
        e1 = linhas["E1"]
        verifica("o titulo diz o dia", "SÁBADO 05/12/2026" in ws["A1"].value, ws["A1"].value)
        verifica("cesta paga, com padrinho e forma",
                 e1[col["Padrinho"]] == padrinho.nome and e1[col["Situação"]] == "Pago"
                 and e1[col["Forma pag."]] == "Pix")
        verifica("festa prometida aparece como promessa",
                 e1[festa] == padrinho.nome and e1[festa + 1].startswith("Promessa"))
        verifica("sem padrinho diz SEM APADRINHAMENTO",
                 linhas["E2"][col["Padrinho"]] == "SEM APADRINHAMENTO")
        verifica("a desistente vem marcada em Saiu?", linhas["E3"][col["Saiu?"]] == "SIM")
        verifica("o cartao de cesta aparece como subido",
                 e1[cabecalho.index("Cartão cesta")] == "Subiu"
                 and e1[cabecalho.index("Cartão festa")] in (None, ""))

        print("\nOs calculos")
        calc = livro["CÁLCULOS"]
        valores = {r[0]: r for r in calc.iter_rows(values_only=True) if r[0]}
        linha = valores[escola.nome]
        # Ônibus, Crianças (sem a desistente), Com cesta, Com festa, Completas, Sem padrinho
        verifica("onibus da instituicao", linha[1] == 2)
        verifica("a desistente nao conta como crianca do dia", linha[2] == 2, str(linha))
        verifica("so o pago conta: 1 cesta, 0 festa", linha[3] == 1 and linha[4] == 0, str(linha))
        verifica("a creche sem dia cai no grupo sem dia", "SEM DIA DEFINIDO" in valores)
        apad = next(r for r in calc.iter_rows(values_only=True) if r[0] == "Apadrinhados (pagos)")
        idx = list(calc.iter_rows(values_only=True)).index(apad)
        total = list(calc.iter_rows(values_only=True))[idx + 1]
        verifica("1 apadrinhamento pago, R$ 130", total[0] == 1 and total[1] == 130, str(total))
        # 3 criancas que vem (E1, E2, S1) x 2 = 6 vagas; so a cesta de E1 esta paga.
        verifica("faltam 5", total[2] == 5, str(total))

        print("\nNada de senha")
        texto = " ".join(
            str(v) for ws in livro for r in ws.iter_rows(values_only=True) for v in r if v
        ).lower()
        verifica("nenhuma coluna de senha", "senha_hash" not in texto and "$argon" not in texto)
    finally:
        limpar(db)
        db.close()

    print(f"\n{ok} ok, {len(falhas)} falha(s)")
    if falhas:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
